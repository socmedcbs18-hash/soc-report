"""Verifikasi PPT hasil generate: baca BALIK file .pptx, ekstrak angka dari tabel dan chart,
lalu bandingkan dengan analytics dan dengan total yang dihitung independen dari raw data."""
import re

import pandas as pd
from pptx import Presentation

from cybersec_report.core.validation import check, overall_status

_SEV = ["Critical", "High", "Medium", "Low"]


def _num(txt):
    """Parse CELL COUNT dari PPT. Count selalu bilangan bulat, jadi semua pemisah ribuan
    (titik locale id / koma locale en) dibuang. Aman untuk kedua locale."""
    t = re.sub(r"[.,\s%]", "", str(txt))
    if not t or t.upper() == "NA" or t.upper() == "N/A":
        return None
    try:
        return float(t)
    except ValueError:
        return None


def _table_rows(tbl):
    return [[c.text.strip() for c in r.cells] for r in tbl.rows]


def extract_numbers(pptx_path):
    """Return {section_name: {...angka dari slide data...}} (berdasarkan nama shape 'TITLE::<section>::data')."""
    prs = Presentation(str(pptx_path))
    out = {}
    for slide in prs.slides:
        title = next((sh for sh in slide.shapes if sh.name.startswith("TITLE::") and sh.name.endswith("::data")), None)
        if title is None:
            continue
        section = title.name.split("::")[1]
        d = {"severity": {}, "attack": {}, "attack_total": {}, "dest": {}, "dest_total": None,
             "country": {}, "country_total": None, "pie": {}, "severity_total": None}
        for sh in slide.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                xml = sh.chart._chartSpace
                cats = [e.text for e in xml.xpath(".//c:ser/c:cat//c:v")]
                vals = [float(e.text) for e in xml.xpath(".//c:ser/c:val//c:v")]
                d["pie"] = dict(zip([c.title() for c in cats], vals))
            if not (getattr(sh, "has_table", False) and sh.has_table):
                continue
            rows = _table_rows(sh.table)
            head, body = rows[0][0], rows[1:]
            if sh.name.startswith("TABLE::severity"):
                for r in body:
                    (d.__setitem__("severity_total", _num(r[2])) if r[0] == "Total" else d["severity"].__setitem__(r[0], _num(r[2])))
            elif "Attack Name" in head:
                sev = re.search(r"\((\w+)\)", head).group(1)
                d["attack"][sev] = {r[0]: _num(r[1]) for r in body if r[0] != "Total"}
                d["attack_total"][sev] = next((_num(r[1]) for r in body if r[0] == "Total"), None)
            elif head == "Dest":
                d["dest"] = {r[0]: _num(r[1]) for r in body if r[0] != "Total"}
                d["dest_total"] = next((_num(r[1]) for r in body if r[0] == "Total"), None)
            elif "Source Country" in head:
                d["country"] = {r[0]: _num(r[1]) for r in body if r[0] != "Total"}
                d["country_total"] = next((_num(r[1]) for r in body if r[0] == "Total"), None)
        out[section] = d
    return out


def verify_ppt(pptx_path, sections):
    nums = extract_numbers(pptx_path)
    checks = []

    def eq(name, a, b):
        ok = a is not None and abs(float(a) - float(b)) < 1e-9
        checks.append(check(name, "PASS" if ok else "FAIL", f"PPT={a} vs analytics={b}", "CRITICAL" if not ok else "INFO", [a, b]))

    for a in sections:
        if a["total_events"] == 0:
            continue
        n = nums.get(a["name"])
        if n is None:
            checks.append(check(f"[{a['name']}] slide data ditemukan", "FAIL", "Slide data tidak ditemukan di PPT", "CRITICAL"))
            continue
        p = f"[{a['name']}] "
        eq(p + "Total severity (PPT) = total event", n["severity_total"], a["total_events"])
        for _, r in a["severity"].iterrows():
            eq(p + f"Severity {r['Severity']}", n["severity"].get(r["Severity"]), r["Count"])
        for sev, df in a["top_by_severity"].items():
            sev_total = float(a["severity"].set_index("Severity").loc[sev, "Count"])
            eq(p + f"Total tabel attack {sev} = total {sev}", n["attack_total"].get(sev), sev_total)
            eq(p + f"Jumlah baris tabel attack {sev} = totalnya", sum(n["attack"].get(sev, {}).values()), sev_total)
        eq(p + "Total destination", n["dest_total"], a["total_events"])
        eq(p + "Jumlah baris destination = total", sum(n["dest"].values()), a["total_events"])
        eq(p + "Total country", n["country_total"], a["total_events"])
        eq(p + "Jumlah baris country = total", sum(n["country"].values()), a["total_events"])
        for _, r in a["severity"].iterrows():
            if r["Count"] > 0:
                eq(p + f"Pie chart {r['Severity']}", n["pie"].get(r["Severity"]), r["Count"])
    return {"checks": checks, "overall": overall_status(checks), "numbers": nums}


def independent_raw_total(df_raw, schema, allowed_actions):
    """Total event dihitung LANGSUNG dari kolom raw (tanpa melewati normalisasi/analytics)."""
    m = schema["mapping"]
    df = df_raw
    if "action" in m:
        df = df[df[m["action"]].astype(str).str.strip().str.lower().isin({a.lower() for a in allowed_actions})]
    if "count" in m:
        return float(pd.to_numeric(df[m["count"]], errors="coerce").clip(lower=0).fillna(0).sum())
    return float(len(df))


def reconcile(raw_total, sections, ppt_numbers):
    """Raw total == Analytics total == PPT total (pada slide data semua section)."""
    analytics_total = float(sum(a["total_events"] for a in sections))
    ppt_total = float(sum((ppt_numbers.get(a["name"], {}).get("severity_total") or 0) for a in sections if a["total_events"] > 0))
    return {"raw_total": raw_total, "analytics_total": analytics_total, "ppt_total": ppt_total,
            "match": abs(raw_total - analytics_total) < 1e-9 and abs(analytics_total - ppt_total) < 1e-9}
