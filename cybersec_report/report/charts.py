"""Chart native python-pptx (editable di PowerPoint, bukan screenshot)."""
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt

PALETTE = ["156082", "E97132", "0F9ED5", "196B24", "A02B93", "4EA72E", "7F7F7F", "C00000"]
SEVERITY_COLORS = {"Critical": "156082", "High": "E97132", "Medium": "0F9ED5", "Low": "196B24",
                   "Info": "7F7F7F", "Unknown": "A6A6A6"}


def _title(chart, title, size=12):
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold = Pt(size), False
        run.font.color.rgb = RGBColor.from_string("595959")
    else:
        chart.has_title = False


def add_pie_chart(slide, categories, values, left, top, width, height, title=None, colors=None, font_pt=9):
    """Pie 2D native. python-pptx tidak dapat menulis pie 3D (deck contoh memakai 3D); 2D dipilih karena
    didukung penuh dan tidak mendistorsi proporsi."""
    cd = CategoryChartData()
    cd.categories = list(categories)
    cd.add_series("Events", [float(v) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Emu(left), Emu(top), Emu(width), Emu(height), cd).chart
    _title(chart, title)
    chart.has_legend = True
    chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    chart.legend.font.size = Pt(font_pt)
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value, dl.number_format, dl.number_format_is_linked = True, "#,##0", False
    dl.font.size = Pt(font_pt)
    colors = colors or [PALETTE[i % len(PALETTE)] for i in range(len(categories))]
    for i, c in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(c)
    return chart


def _cartesian(slide, kind, categories, series, left, top, width, height, title, colors, number_format):
    cd = CategoryChartData()
    cd.categories = list(categories)
    for name, vals in series.items():
        cd.add_series(name, [float(v) for v in vals])
    chart = slide.shapes.add_chart(kind, Emu(left), Emu(top), Emu(width), Emu(height), cd).chart
    _title(chart, title)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        chart.legend.font.size = Pt(9)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = number_format, False
    plot.data_labels.font.size = Pt(9)
    for i, s in enumerate(plot.series):
        col = RGBColor.from_string((colors or PALETTE)[i % len(colors or PALETTE)])
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            s.format.line.color.rgb = col
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = col
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.has_major_gridlines = False
    return chart


def add_bar_chart(slide, categories, series, left, top, width, height, title=None, horizontal=True, colors=None, number_format="#,##0"):
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = _cartesian(slide, kind, categories, series, left, top, width, height, title, colors, number_format)
    if horizontal:
        chart.category_axis.reverse_order = True
    return chart


def add_line_chart(slide, categories, series, left, top, width, height, title=None, colors=None, number_format="#,##0"):
    return _cartesian(slide, XL_CHART_TYPE.LINE_MARKERS, categories, series, left, top, width, height, title, colors, number_format)
