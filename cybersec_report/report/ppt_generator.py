"""Generator PPT IPS. Template dasar = deck contoh (theme tetap) + background gambar hasil ekstraksi.
Semua angka datang dari analytics (yang dihitung dari raw data); tidak ada angka hardcode."""
import math

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from cybersec_report.core.config_loader import ASSET_DIR, TEMPLATE_DIR, load_report_structure
from cybersec_report.core.fmt import fmt_int, fmt_pct
from cybersec_report.report import slide_components as SC
from cybersec_report.report.charts import SEVERITY_COLORS
from cybersec_report.report.narrative import ips_narrative

TEMPLATE = TEMPLATE_DIR / "CyberSec_Monthly_Template.pptx"
COL_X = (3860000, 8000000)
COL_W = 3850000


def _blank_layout(prs):
    return next((l for l in prs.slide_layouts if l.name == "Blank"), prs.slide_layouts[6])


def _slide(prs, bg):
    s = prs.slides.add_slide(_blank_layout(prs))
    SC.set_background(s, ASSET_DIR / bg)
    return s


def _tb(slide, left, top, w, h):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(w), Emu(h))
    tb.text_frame.word_wrap = True
    return tb


def build_cover(prs, period, source_files, section_titles):
    s = _slide(prs, "bg_cover.png")
    tb = _tb(s, 438049, 695328, 8806316, 523220)
    SC._run(tb.text_frame.paragraphs[0], "Cybersecurity  Monthly Deck", 28, True, "156082", "Aptos")
    tb2 = _tb(s, 491516, 1502435, 7000000, 3500000)
    p = tb2.text_frame.paragraphs[0]
    SC._run(p, f"Periode: {period['month_name']} {period['report_year']}", 16, True, "0E2841")
    for i, t in enumerate(section_titles, 1):
        q = tb2.text_frame.add_paragraph()
        q.space_before = Pt(8)
        SC._run(q, f"{i}. {t}", 16, True, "0E2841")
    q = tb2.text_frame.add_paragraph()
    q.space_before = Pt(14)
    SC._run(q, "Sumber data: " + ", ".join(source_files), 11, False, "595959")
    return s


def build_section_header(prs, title, period, actions):
    s = _slide(prs, "bg_section.png")
    s.shapes.add_picture(str(ASSET_DIR / "section_art.png"), Emu(3939988), Emu(1371600), Emu(4861111), Emu(3251947))
    sd, ed = period["start_date"], period["end_date"]
    lines = [[(title, False)],
             [("Date Range: ", False), (f"{sd:%Y/%m/%d} - {ed:%Y/%m/%d}", True)],
             [("Action: ", False), (", ".join(actions) if actions else "semua action", True)]]
    for i, segs in enumerate(lines):
        tb = _tb(s, 272725, 5354482 + i * 347000, 9000000, 369332)
        for t, b in segs:
            SC._run(tb.text_frame.paragraphs[0], t, 18, b, "000000", "Aptos")
    return s


def _data_slide_tables(a, settings, loc):
    I, P = (lambda x: fmt_int(x, loc)), (lambda x: fmt_pct(x, 2, loc))
    n = settings["top_n"]
    out = []

    def mk(header, df, label, widths=(2400000, 600000, 850000)):
        rows = [[r[label], I(r["Count"]), P(r["Percent"])] for _, r in df.iterrows()]
        rows.append(["Total", I(df["Count"].sum()), P(100.0)])
        return {"header": header, "rows": rows, "widths": widths}

    for sev, df in a["top_by_severity"].items():
        has_others = (df["Attack Name"] == "Others").any()
        hdr = SC.severity_header(f"Top {n} Attack Name" if has_others else "Attack Name", sev)
        hdr = [hdr, "Count", "By%"]
        out.append(mk(hdr, df, "Attack Name"))
    out.append(mk(["Dest", "Count", "By%"], a["top_destination"], "Dest", (2450000, 550000, 850000)))
    c = a["top_country"]
    out.append(mk([f"Top {n} Source Country" if (c["Source Country"] == "Others").any() else "Source Country", "Count", "By%"], c, "Source Country", (2450000, 550000, 850000)))
    return out


def _masonry(tables, font_pt, gap=110000):
    tops, placed = [SC.CONTENT_TOP, SC.CONTENT_TOP], []
    for t in tables:
        h = SC.estimate_table_height(t["header"], t["rows"], t["widths"], font_pt)
        col = 0 if tops[0] <= tops[1] else 1
        placed.append((t, col, tops[col], h))
        tops[col] += h + gap
    return placed, max(tops) - gap


def build_data_slide(prs, a, period, settings, struct, source_file):
    loc = settings.get("number_locale", "id")
    I, P = (lambda x: fmt_int(x, loc)), (lambda x: fmt_pct(x, 2, loc))
    env = a.get("label") or a["environment"] or a["name"]
    s = _slide(prs, "bg_content.png")
    SC.add_title(s, struct["title_pattern"].format(env=env, month=period["month_name"], year=period["report_year"]),
                 name=f"TITLE::{a['name']}::data")
    SC.add_subtitle(s, struct["purpose"], "memastikan jalur komunikasi antar server/klien aman sebelum ke aplikasi.")
    SC.add_logo(s, ASSET_DIR / "logo_fortinet.png", 9826718, -134471, 1704975, 1143000)
    sev = a["severity"]
    shown = sev[(sev["Count"] > 0)]
    SC.add_pie_chart(s, shown["Severity"].tolist(), shown["Count"].tolist(), 160000, 1100000, 3550000, 2650000,
                     title=f"Intrusion Prevention System - Firewall {env}",
                     colors=[SEVERITY_COLORS.get(x, "7F7F7F") for x in shown["Severity"]])
    rows = [[r["Severity"], r["Description"], I(r["Count"])] for _, r in sev.iterrows()] + [["Total", "", I(sev["Count"].sum())]]
    SC.add_table(s, ["Severity", "Description", "Count"], rows, 160000, 3820000, [760000, 2060000, 730000], font_pt=8,
                 align=["l", "l", "r"], name=f"TABLE::severity::{a['name']}")
    tables = _data_slide_tables(a, settings, loc)
    for fp in (9, 8, 7):
        placed, bottom = _masonry(tables, fp)
        if bottom <= SC.CONTENT_BOTTOM:
            break
    for t, col, top, _ in placed:
        SC.add_table(s, t["header"], t["rows"], COL_X[col], top, t["widths"], font_pt=fp, name="TABLE::" + (t["header"][0] if isinstance(t["header"][0], str) else "".join(x for x, _ in t["header"][0])))
    SC.add_footer(s, f"Sumber: {source_file} · Total event = SUM(count) dari {I(a['reconciliation']['raw_records_in_file'])} record · % tabel attack terhadap total severity")
    return s


def build_analysis_slide(prs, a, period, narrative, struct):
    env = a.get("label") or a["environment"] or a["name"]
    s = _slide(prs, "bg_content.png")
    SC.add_title(s, struct["title_pattern"].format(env=env, month=period["month_name"], year=period["report_year"]),
                 name=f"TITLE::{a['name']}::analysis")
    SC.add_subtitle(s, struct["purpose"], "memastikan jalur komunikasi antar server/klien aman sebelum ke aplikasi.")
    SC.add_logo(s, ASSET_DIR / "logo_fortinet.png", 9826718, -134471, 1704975, 1143000)
    SC.add_conclusion(s, narrative["kesimpulan"], 100000, 1100000, 4050000, 5200000)
    SC.add_stacked_blocks(s, [("Implikasi", narrative["implikasi"], "BLOCK::implikasi"),
                              ("Insight", narrative["insight"], "BLOCK::insight")], 4250000, 1100000, 3750000, 5200000)
    SC.add_stacked_blocks(s, [("Rekomendasi", narrative["rekomendasi"], "BLOCK::rekomendasi"),
                              ("Tindakan & Pencegahan", narrative["tindakan"] + narrative["pencegahan"], "BLOCK::tindakan")],
                          8100000, 1100000, 3900000, 5200000)
    SC.add_footer(s, "Narasi otomatis (rule-based) dari data; menyatakan indikasi, bukan kesimpulan compromise.")
    return s


def build_source_ip_slide(prs, a, period, profile, rep_summary, rep_stats, settings, struct, source_file):
    loc = settings.get("number_locale", "id")
    I = lambda x: fmt_int(x, loc)
    env = a.get("label") or a["environment"] or a["name"]
    s = _slide(prs, "bg_content.png")
    SC.add_title(s, f"(IPS) Top Source IP – {env} – {period['month_name']} {period['report_year']}", name=f"TITLE::{a['name']}::source_ip")
    SC.add_subtitle(s, "Profil source IP berdasarkan jumlah event, variasi signature, dan reputasi (Internal Normalized Threat Score).",
                    "mengidentifikasi sumber serangan berulang dan kandidat tindak lanjut.")
    SC.add_logo(s, ASSET_DIR / "logo_fortinet.png", 9826718, -134471, 1704975, 1143000)
    assessed = rep_summary["assessed"] if rep_summary else 0
    hi = rep_summary["high_or_critical"] if rep_summary else 0
    rep_val = f"{I(assessed)}" if assessed else "Not run"
    rep_sub = (f"dari {I(rep_stats['public_queued'])} IP publik diantrikan" if rep_stats else "provider belum dijalankan/tidak tersedia")
    cards = [("Total Events", I(a["total_events"]), f"{I(a['reconciliation']['raw_records_in_file'])} raw records"),
             ("Unique Source IP", I(a["unique"]["source_ip"]), None),
             ("Unique Attack Signature", I(a["unique"]["attack"]), None),
             ("IP Reputation Assessed", rep_val, rep_sub),
             ("IP Risk High/Critical", I(hi) if assessed else "N/A", "Internal Normalized Threat Score" if assessed else None)]
    w, gap = 2250000, 110000
    for i, (lab, val, sub) in enumerate(cards):
        SC.add_kpi_card(s, 160000 + i * (w + gap), 1150000, w, 880000, lab, val, sub)
    top = profile.head(10)
    na = lambda v, f=lambda x: x: "N/A" if pd.isna(v) else f(v)
    rows = [[r["Source IP"], r["Country"], I(r["Attack Count"]), I(r["Unique Attack"]), I(r["Unique Destination"]), I(r["Critical"]), I(r["High"]),
             na(r["VT Malicious"], lambda x: I(x)), na(r["AbuseIPDB Score"], lambda x: I(x)),
             na(r["Threat Score"], lambda x: f"{x:.1f}".replace(".", ",") if loc == "id" else f"{x:.1f}"), r["Risk Level"]] for _, r in top.iterrows()]
    hdr = ["Source IP", "Country", "Attack Count", "Unique Attack", "Unique Dest", "Critical", "High", "VT Malicious", "AbuseIPDB Score", "Threat Score", "Risk Level"]
    widths = [1450000, 1400000, 950000, 950000, 950000, 800000, 700000, 1000000, 1100000, 1000000, 1050000]
    shp, h = SC.add_table(s, hdr, rows, 160000, 2200000, widths, font_pt=9, align=["l", "l"] + ["r"] * 8 + ["c"], bold_last_row=False,
                          name=f"TABLE::source_ip::{a['name']}")
    note = ("Internal Normalized Threat Score (0–100) adalah skor internal hasil normalisasi data VirusTotal dan AbuseIPDB (bobot di config/scoring.json), "
            "BUKAN skor resmi vendor. 'N/A' = provider tidak tersedia atau IP tidak masuk antrian cek (pengecekan diprioritaskan ke IP dengan Critical dan event terbanyak). "
            "Country berasal dari log (GeoIP) dan bukan atribusi pelaku. Urutan: Threat Score lalu Attack Count.")
    nb = _tb(s, 160000, 2200000 + h + 120000, 11700000, 700000)
    SC._run(nb.text_frame.paragraphs[0], note, 9, False, "595959", italic=True)
    SC.add_footer(s, f"Sumber: {source_file} + provider threat intelligence yang aktif")
    return s


def generate_ips_deck(sections, period, source_file, findings_by_section, rep_by_section, settings, output_path=None, template=None):
    """sections: output compute_ips_sections. rep_by_section[name] = {'profile','summary','stats'}."""
    struct = next(x for x in load_report_structure()["sections"] if x["id"] == "ips")
    prs = Presentation(str(template or TEMPLATE))
    titles = [x["name"] for x in sections]
    build_cover(prs, period, [source_file], ["Intrusion Prevention System – " + ", ".join(s.get("label") or s["environment"] or "All Devices" for s in sections)])
    actions = sorted({x for s in sections for x in (s.get("actions") or [])})
    build_section_header(prs, "Intrusion Prevention System – " + ", ".join(s.get("label") or s["environment"] or "All Devices" for s in sections), period, actions)
    for a in sections:
        if a["total_events"] == 0:
            s = _slide(prs, "bg_content.png")
            SC.add_title(s, f"(IPS) {a.get('label') or a['environment']} – {period['month_name']} {period['report_year']}", name=f"TITLE::{a['name']}::data")
            tb = _tb(s, 700000, 2500000, 10000000, 800000)
            SC._run(tb.text_frame.paragraphs[0], "Tidak ada event serangan pada scope data periode ini.", 20, True, "0E2841")
            continue
        findings = findings_by_section[a["name"]]
        narr = ips_narrative(findings)
        rep = rep_by_section[a["name"]]
        build_data_slide(prs, a, period, settings, struct, source_file)
        build_analysis_slide(prs, a, period, narr, struct)
        build_source_ip_slide(prs, a, period, rep["profile"], rep["summary"], rep["stats"], settings, struct, source_file)
    if output_path:
        prs.save(str(output_path))
    return prs
