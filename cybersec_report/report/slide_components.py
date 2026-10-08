"""Komponen slide reusable. Gaya mengikuti deck contoh: header biru (dari background template),
tabel header 156082 + baris C0E6F5 bergantian, font Aptos (theme)."""
import math

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

from cybersec_report.report.charts import add_bar_chart, add_line_chart, add_pie_chart  # noqa: F401 (re-export)

HEADER_FILL, BAND_FILL, WHITE, DARK = "156082", "C0E6F5", "FFFFFF", "0E2841"
SEV_TEXT = {"Critical": "FF0000", "High": "FFC000", "Medium": "FFFF00", "Low": "00B050"}
SLIDE_W, SLIDE_H = 12192000, 6858000
CONTENT_TOP, CONTENT_BOTTOM = 1100000, 6330000
_NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def set_background(slide, image_path):
    """Template sample memakai GAMBAR background per slide (bukan layout/master)."""
    _, rid = slide.part.get_or_add_image_part(str(image_path))
    bg = parse_xml(f'<p:bg {_NS}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/>'
                   f'<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    cSld = slide._element.cSld
    for old in cSld.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}bg"):
        cSld.remove(old)
    cSld.insert(0, bg)


def _text(slide, left, top, width, height, name=None):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    if name:
        tb.name = name
    tb.text_frame.word_wrap = True
    return tb


def _run(p, text, size, bold=False, color=DARK, font=None, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
    r.font.color.rgb = RGBColor.from_string(color)
    if font:
        r.font.name = font
    return r


def add_title(slide, text, name=None, size=None):
    """Ukuran font adaptif terhadap panjang judul supaya tidak wrap menabrak subtitle."""
    if size is None:
        size = 24 if len(text) <= 30 else 20 if len(text) <= 40 else 17
    tb = _text(slide, 2300000, 30000, 7450000, 480000, name)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, text, size, True, WHITE, "Lato Black")
    return tb


def add_subtitle(slide, text, goal=None):
    tb = _text(slide, 2300000, 500000, 7450000, 560000)
    p = tb.text_frame.paragraphs[0]
    _run(p, text, 9, True, WHITE)
    if goal:
        p2 = tb.text_frame.add_paragraph()
        _run(p2, "Tujuan: ", 9, True, WHITE)
        _run(p2, goal, 9, False, WHITE)
    return tb


def add_footer(slide, text):
    tb = _text(slide, 2400000, 6440000, 7400000, 250000)
    _run(tb.text_frame.paragraphs[0], text, 8, False, "595959", italic=True)
    return tb


def add_logo(slide, image_path, left, top, width, height):
    return slide.shapes.add_picture(str(image_path), Emu(left), Emu(top), Emu(width), Emu(height))


# ---------- tabel ----------
def _chars_per_line(width_emu, font_pt):
    return max(4, int((width_emu / 12700 - 8) / (font_pt * 0.5)))


def estimate_row_heights(rows, col_widths, font_pt, min_h):
    line_h = int(font_pt * 1.25 * 12700)
    heights = []
    for row in rows:
        lines = max(math.ceil(len(str(c)) / _chars_per_line(w, font_pt)) if str(c) else 1 for c, w in zip(row, col_widths))
        heights.append(max(min_h, lines * line_h + 60000))
    return heights


def estimate_table_height(header, rows, col_widths, font_pt=9, min_h=190000):
    hdr = [h if isinstance(h, str) else "".join(t for t, _ in h) for h in header]
    return sum(estimate_row_heights([hdr] + rows, col_widths, font_pt, min_h))


def add_table(slide, header, rows, left, top, col_widths, font_pt=9, min_row_h=190000,
              align=None, bold_last_row=True, name=None):
    """header: list str | list[(teks, warna)] (untuk kata severity berwarna). rows: list[list[str]]."""
    ncols, nrows = len(col_widths), len(rows) + 1
    hdr_txt = [h if isinstance(h, str) else "".join(t for t, _ in h) for h in header]
    heights = estimate_row_heights([hdr_txt] + rows, col_widths, font_pt, min_row_h)
    shape = slide.shapes.add_table(nrows, ncols, Emu(left), Emu(top), Emu(sum(col_widths)), Emu(sum(heights)))
    if name:
        shape.name = name
    tbl = shape.table
    tbl.first_row, tbl.horz_banding = True, False
    align = align or (["l"] + ["r"] * (ncols - 1))
    amap = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = Emu(w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = Emu(h)

    def fill(cell, hexcolor):
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(hexcolor)

    for i in range(nrows):
        is_hdr, is_total = i == 0, (bold_last_row and i == nrows - 1)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Emu(50000)
            cell.margin_top = cell.margin_bottom = Emu(20000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill(cell, HEADER_FILL if is_hdr else (BAND_FILL if (i % 2 == 1 or is_total) else WHITE))
            p = cell.text_frame.paragraphs[0]
            p.alignment = amap[align[j]] if not is_hdr else (PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
            if is_hdr:
                segs = header[j] if not isinstance(header[j], str) else [(header[j], WHITE)]
                for t, col in segs:
                    _run(p, t, font_pt, True, col)
            else:
                _run(p, str(rows[i - 1][j]), font_pt, is_total, "000000")
    return shape, sum(heights)


def severity_header(prefix, severity, suffix=""):
    """'Top 5 Attack Name (' + Critical berwarna + ')' seperti pada deck contoh."""
    return [(f"{prefix} (", WHITE), (severity, SEV_TEXT.get(severity, WHITE)), (")" + suffix, WHITE)]


# ---------- teks naratif ----------
def estimate_text_lines(paragraphs, width_emu, font_pt):
    cpl = _chars_per_line(width_emu, font_pt)
    return sum(max(1, math.ceil(len(p) / cpl)) for p in paragraphs)


def fit_font(paragraphs, width_emu, height_emu, sizes=(11, 10.5, 10, 9.5, 9, 8.5, 8)):
    for s in sizes:
        need = estimate_text_lines(paragraphs, width_emu - 150000, s) * s * 1.3 * 12700 + len(paragraphs) * s * 0.5 * 12700
        if need <= height_emu:
            return s
    return sizes[-1]


def add_text_block(slide, heading, bullets, left, top, width, height, sizes=(11, 10.5, 10, 9.5, 9, 8.5, 8), name=None):
    sub = lambda b: b.startswith("– ")
    paras = [heading] + [("      " + b if sub(b) else "• " + b) for b in bullets]
    size = fit_font(paras, width, height, sizes)
    tb = _text(slide, left, top, width, height, name)
    p = tb.text_frame.paragraphs[0]
    _run(p, heading, size + 1, True, DARK)
    for b in bullets:
        q = tb.text_frame.add_paragraph()
        q.space_before = Pt(1 if sub(b) else 4)
        _run(q, ("      " + b) if sub(b) else "• " + b, size, False, "000000")
    return tb, size


def add_stacked_blocks(slide, blocks, left, top, width, total_height, gap=120000):
    """blocks: [(heading, bullets, shape_name)]. Tinggi tiap blok proporsional terhadap perkiraan kebutuhan
    teksnya, jadi blok pendek tidak menyisakan ruang kosong dan blok panjang tidak terpaksa berfont kecil."""
    needs = [estimate_text_lines([h] + ["• " + x for x in b], width - 150000, 10) + 0.6 * len(b) + 1.5 for h, b, _ in blocks]
    avail, y = total_height - gap * (len(blocks) - 1), top
    for (h, b, n), need in zip(blocks, needs):
        bh = int(avail * need / sum(needs))
        add_text_block(slide, h, b, left, y, width, bh, name=n)
        y += bh + gap


def add_conclusion(slide, bullets, left, top, width, height, heading="Kesimpulan – Firewall / IPS Monitoring"):
    return add_text_block(slide, heading, bullets, left, top, width, height, name="BLOCK::kesimpulan")


def add_recommendation(slide, bullets, left, top, width, height, heading="Insight & Rekomendasi"):
    return add_text_block(slide, heading, bullets, left, top, width, height, name="BLOCK::rekomendasi")


def add_action_plan(slide, bullets, left, top, width, height, heading="Tindakan & Pencegahan"):
    return add_text_block(slide, heading, bullets, left, top, width, height, name="BLOCK::tindakan")


def add_kpi_card(slide, left, top, width, height, label, value, sub=None, accent=HEADER_FILL):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    shp.name = f"KPI::{label}"
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string("EAF1FA")
    shp.line.color.rgb = RGBColor.from_string(accent)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Emu(60000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, label, 9, False, "595959")
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    _run(p2, str(value), 20, True, accent)
    if sub:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        _run(p3, sub, 8, False, "595959")
    return shp
