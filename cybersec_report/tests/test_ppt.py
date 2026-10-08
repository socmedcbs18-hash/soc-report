import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from cybersec_report.report import ppt_verify
from cybersec_report.report import slide_components as SC
from cybersec_report.tests._helpers import make_pipeline


class TestPPT(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pl = make_pipeline()
        cls.out = Path(tempfile.mkdtemp())
        cls.res = cls.pl.generate(out_dir=cls.out, save_history=False)

    def test_ppt_generation(self):
        for key in ("pptx", "summary", "ioc", "audit", "validation"):
            self.assertTrue(Path(self.res["files"][key]).exists(), key)
        self.assertRegex(Path(self.res["files"]["pptx"]).name, r"^Cybersecurity_Monthly_Report_2026_09\.pptx$")
        prs = Presentation(self.res["files"]["pptx"])
        self.assertEqual(len(prs.slides), 5)  # cover, section, data, analisis, source IP

    def test_raw_equals_analytics_equals_ppt(self):
        r = self.res["reconciliation"]
        self.assertEqual((r["raw_total"], r["analytics_total"], r["ppt_total"]), (939.0, 939.0, 939.0))
        self.assertTrue(r["match"])
        self.assertEqual(self.res["ppt_verification"]["overall"], "PASS")

    def test_verifier_detects_tampered_number(self):
        prs = Presentation(self.res["files"]["pptx"])
        for sl in prs.slides:
            for sh in sl.shapes:
                if getattr(sh, "has_table", False) and sh.has_table and sh.name.startswith("TABLE::severity"):
                    cell = sh.table.cell(1, 2)  # baris Critical
                    cell.text_frame.paragraphs[0].runs[0].text = "999"
        bad = self.out / "tampered.pptx"
        prs.save(bad)
        res = ppt_verify.verify_ppt(bad, self.pl.sections)
        self.assertEqual(res["overall"], "FAIL")

    def test_slides_not_hardcoded_month(self):
        txt = " ".join(sh.text_frame.text for sl in Presentation(self.res["files"]["pptx"]).slides
                       for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("September 2026", txt)
        for stale in ("April 2026", "Maret 2026", "Desember 2025"):
            self.assertNotIn(stale, txt)

    def test_narrative_hedged_no_compromise_claim(self):
        txt = " ".join(sh.text_frame.text for sl in Presentation(self.res["files"]["pptx"]).slides
                       for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("tidak cukup untuk membuktikan", txt)
        for forbidden in ("tanpa insiden breach", "attacker failed", "pasti tidak ada compromise"):
            self.assertNotIn(forbidden, txt)

    def test_mom_only_when_previous_month_exists(self):
        self.assertIsNone(self.res["mom"]["IPS (All Devices)"])

    def test_components_smoke(self):
        from pptx import Presentation as P
        prs = P()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        SC.add_title(s, "Judul")
        SC.add_kpi_card(s, 0, 0, 1000000, 800000, "KPI", "1")
        SC.add_bar_chart(s, ["a", "b"], {"x": [1, 2]}, 0, 900000, 3000000, 2000000)
        SC.add_line_chart(s, ["m1", "m2"], {"x": [1, 2]}, 3100000, 900000, 3000000, 2000000)
        SC.add_pie_chart(s, ["a", "b"], [1, 2], 6200000, 900000, 3000000, 2000000)
        SC.add_table(s, ["h1", "h2"], [["a", "1"], ["Total", "1"]], 0, 3000000, [1000000, 500000])
        SC.add_footer(s, "f")
        self.assertEqual(len(s.shapes), 7)  # title, kpi, bar, line, pie, table, footer

    def test_estimate_table_height_grows_with_long_text(self):
        short = SC.estimate_table_height(["a"], [["x"]], [1000000])
        long = SC.estimate_table_height(["a"], [["x" * 200]], [1000000])
        self.assertGreater(long, short)


if __name__ == "__main__":
    unittest.main()
