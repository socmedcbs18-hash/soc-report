"""Test sumber Akamai: parser (layout berdampingan & bertumpuk, zero-width char), validasi, analytics, snapshot, compare, UI (stub).
Data uji SINTETIS (bukan data ALTO); hanya meniru tata letak file dari tim Akamai."""
import io
import tempfile
import unittest
from pathlib import Path

import openpyxl

from cybersec_report.sources import akamai as ak
from cybersec_report.sources import akamai_store as store
from cybersec_report.tests.test_ui_smoke import FakeUpload, StubStreamlit, UIHarness

ZW = "​"


def _side_by_side(ws, cat_hdr, cats, dests, ctry, bad_country_total=None):
    cols = [(0, cat_hdr, cats), (4, "Dest", dests), (8, "Country", ctry)]
    for c0, hdr, rows in cols:
        ws.cell(1, c0 + 1, hdr + ZW); ws.cell(1, c0 + 2, "Total Hits" + ZW); ws.cell(1, c0 + 3, "By %" + ZW)
        for i, (lab, v) in enumerate(rows, start=2):
            ws.cell(i, c0 + 1, lab); ws.cell(i, c0 + 2, v)
        tot = sum(v for _, v in rows)
        ws.cell(len(rows) + 2, c0 + 1, "Total" + ZW)
        ws.cell(len(rows) + 2, c0 + 2, bad_country_total if (hdr == "Country" and bad_country_total) else tot)


def _stacked(ws, cats, dests, ctry):
    r = 1
    for hdr, rows in (("Category", cats), ("Dest", dests), ("Country", ctry)):
        ws.cell(r, 1, hdr); ws.cell(r, 2, "Total Hits" + ZW); ws.cell(r, 3, "By %" + ZW)
        for lab, v in rows:
            r += 1
            ws.cell(r, 1, lab); ws.cell(r, 2, v)
        r += 1
        ws.cell(r, 1, "Total" + ZW); ws.cell(r, 2, sum(v for _, v in rows))
        r += 3


def make_book(scale=1, waf_country_total=None, drop=()):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    s = lambda rows: [(l, v * scale) for l, v in rows]
    if "WAF Attack" not in drop:
        ws = wb.create_sheet("WAF Attack")
        _side_by_side(ws, "Attack Type", s([("Local File Inclusion", 400), ("Custom Rules", 300), ("Others", 100)]),
                      s([("www.a.id", 500), ("www.b.id", 300)]), s([("US", 450), ("FR", 350)]),
                      bad_country_total=waf_country_total)
    ws = wb.create_sheet("BOT Attack")
    _side_by_side(ws, "BOT Type", s([("Browser Impersonator", 700), ("Declared Bots", 300)]),
                  s([("www.a.id", 900), ("www.c.id", 100)]), s([("US", 600), ("ID", 400)]))
    ws = wb.create_sheet("Bad Reputation")
    _stacked(ws, s([("Web Attackers", 90), ("Scanning Tools", 10)]), s([("www.a.com", 60), ("www.a.id", 40)]), s([("FR", 70), ("Others", 30)]))
    ws = wb.create_sheet("DOS")
    _stacked(ws, s([("Network Control", 8000), ("Rate Control", 2000)]), s([("qris.x.id", 7000), ("www.a.id", 3000)]), s([("US", 6000), ("ID", 4000)]))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


class TestParse(unittest.TestCase):
    def test_parses_both_layouts_and_ignores_zero_width(self):
        p = ak.parse_workbook(make_book(), "Report_WAF_Akamai_September_2026.xlsx")
        self.assertEqual(set(p["sections"]), {"waf", "bot", "badrep", "dos"})
        s = ak.summarize(p)
        self.assertEqual(s["totals"], {"badrep": 100, "waf": 800, "bot": 1000, "dos": 10000})
        self.assertEqual(s["grand_total"], 11900)
        self.assertAlmostEqual(s["shares"]["dos"], 10000 / 11900 * 100)
        self.assertFalse([c for c in p["checks"] if c["level"] == "error"])
        self.assertEqual(ak.table_rows(p["sections"]["waf"], "category")[-1][0], "Others")

    def test_total_mismatch_is_flagged_not_hidden(self):
        p = ak.parse_workbook(make_book(waf_country_total=810), "x.xlsx")
        errs = [c["msg"] for c in p["checks"] if c["level"] == "error"]
        self.assertEqual(len(errs), 1)
        self.assertIn("800", errs[0]); self.assertIn("810", errs[0])

    def test_missing_sheet_warns_and_continues(self):
        p = ak.parse_workbook(make_book(drop=("WAF Attack",)), "x.xlsx")
        self.assertNotIn("waf", p["sections"])
        self.assertTrue(any("waf attack" in c["msg"] for c in p["checks"] if c["level"] == "warn"))
        self.assertNotIn("waf", ak.summarize(p)["totals"])

    def test_unknown_or_broken_file_raises_clean_error(self):
        with self.assertRaises(ak.AkamaiError):
            ak.parse_workbook(io.BytesIO(b"not an xlsx"), "x.xlsx")
        wb = openpyxl.Workbook(); wb.active.title = "Lain"
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        with self.assertRaises(ak.AkamaiError):
            ak.parse_workbook(buf, "x.xlsx")

    def test_text_numbers_are_rejected_not_guessed(self):
        wb = openpyxl.load_workbook(make_book())
        wb["DOS"]["B2"] = "1.234"
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        p = ak.parse_workbook(buf, "x.xlsx")
        self.assertTrue(any("bukan angka" in c["msg"] for c in p["checks"]))


class TestInsights(unittest.TestCase):
    def setUp(self):
        self.p = ak.parse_workbook(make_book(), "x.xlsx")
        self.ins = ak.build_insights(self.p)

    def test_hedged_and_grounded(self):
        text = " ".join(i["text"] for i in self.ins).lower()
        import re
        for bad in (r"\bpasti\b", "berhasil dibobol", "kompromi"):
            self.assertIsNone(re.search(bad, text), bad)
        self.assertIn("tidak dapat dipastikan semuanya diblokir", text)
        titles = " ".join(i["title"] for i in self.ins)
        self.assertIn("DoS Traffic mendominasi", titles)
        self.assertIn("Host yang muncul di banyak kategori", titles)

    def test_similar_hostnames_flagged_not_merged(self):
        self.assertTrue(any("domain berbeda" in i["title"] and "www.a.com" in i["text"] and "www.a.id" in i["text"] for i in self.ins))
        hosts = {h["host"] for h in ak.summarize(self.p)["hosts"]}
        self.assertIn("www.a.com", hosts); self.assertIn("www.a.id", hosts)

    def test_excel_export(self):
        path = Path(tempfile.mkdtemp()) / "a.xlsx"
        ak.to_excel(self.p, path)
        wb = openpyxl.load_workbook(path)
        self.assertIn("Ringkasan", wb.sheetnames); self.assertIn("Validasi", wb.sheetnames)


class TestStore(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def test_save_load_list_delete(self):
        p = ak.parse_workbook(make_book(), "x.xlsx")
        path = store.save_snapshot(p, 9, 2026, self.d)
        store.save_snapshot(p, 9, 2026, self.d)             # simpan ulang = timpa
        self.assertEqual(len(store.list_snapshots(self.d)), 1)
        self.assertEqual(store.list_snapshots(self.d)[0]["label"], "AKAMAI-September-2026")
        back = store.load_snapshot(path)
        self.assertEqual(ak.summarize(back)["grand_total"], 11900)
        self.assertTrue(store.delete_snapshot(path, self.d))
        self.assertEqual(store.list_snapshots(self.d), [])

    def test_delete_refuses_other_paths(self):
        f = Path(tempfile.mkdtemp()) / "AKAMAI_2026-09.json"
        f.write_text("{}")
        with self.assertRaises(ValueError):
            store.delete_snapshot(f, self.d)
        self.assertTrue(f.exists())

    def test_compare_separates_missing_labels_from_zero(self):
        a = ak.parse_workbook(make_book(scale=1), "a.xlsx")
        b = ak.parse_workbook(make_book(scale=2), "b.xlsx")
        b["sections"]["waf"]["tables"]["destination"]["items"][1]["label"] = "www.new.id"
        r = store.compare(a, b)
        self.assertEqual(r["grand"], (11900, 23800))
        dest = r["sections"]["waf"]["destination"]
        self.assertEqual(dest["only_a"], ["www.b.id"]); self.assertEqual(dest["only_b"], ["www.new.id"])
        self.assertEqual(dest["both"], [("www.a.id", 500, 1000)])


class TestAkamaiUI(UIHarness, unittest.TestCase):
    def _up(self, name="Report_WAF_Akamai_September_2026.xlsx", **kw):
        up = FakeUpload(make_book(**kw).getvalue())
        up.name = name
        return up

    def test_menu_is_live_not_placeholder_and_empty_state_ok(self):
        stub = StubStreamlit(choices={"cs_source": "akamai"})
        self._run(stub)
        self.assertNotIn("error", stub.calls)
        self.assertFalse(stub.images)

    def test_upload_renders_dashboard(self):
        stub = StubStreamlit(uploads={"cs_ak_upload": self._up()}, choices={"cs_source": "akamai"})
        self._run(stub)
        html = "".join(stub.html)
        for needle in ("Akamai WAF / Anti-DDoS / Anti-BOT · ", "Traffic by attack type", "Data validation: PASS", "DoS Traffic", "qris.x.id", "Local File Inclusion", "11.900"):
            self.assertIn(needle, html)
        self.assertNotIn("error", stub.calls)
        self.assertNotIn("Attack Name", html)

    def test_save_requires_period_confirmation_then_compare(self):
        s1 = StubStreamlit(uploads={"cs_ak_upload": self._up()}, clicks={"cs_ak_save"}, choices={"cs_source": "akamai"})
        self._run(s1)
        self.assertEqual(store.list_snapshots(), [])               # belum konfirmasi periode -> tombol non-aktif secara logika UI
        s2 = StubStreamlit(uploads={"cs_ak_upload": self._up()}, clicks={"cs_ak_save"}, choices={"cs_source": "akamai", "cs_ak_period_ok": True})
        self._run(s2)
        snaps = store.list_snapshots()
        self.assertEqual([x["label"] for x in snaps], ["AKAMAI-September-2026"])
        s3 = StubStreamlit(uploads={"cs_ak_upload": self._up("Report_WAF_Akamai_Oktober_2026.xlsx", scale=2)}, clicks={"cs_ak_save"},
                           choices={"cs_source": "akamai", "cs_ak_period_ok": True})
        self._run(s3)
        self.assertEqual(len(store.list_snapshots()), 2)
        s4 = StubStreamlit(choices={"cs_source": "akamai"})
        self._run(s4)
        self.assertIn("→", "".join(s4.html))
        self.assertNotIn("error", s4.calls)

    def test_ips_upload_untouched_by_akamai_page_and_same_tab_pattern_as_ips(self):
        stub = StubStreamlit(uploads={"cs_ak_upload": self._up()}, choices={"cs_source": "akamai"})
        self._run(stub)
        self.assertNotIn("exception", stub.calls)
        self.assertTrue(stub.session_state.get("cs_ak_parsed"))
        self.assertNotIn("cs_pl", stub.session_state)        # pipeline IPS tidak dibuat dari file Akamai

    def test_tables_tab_has_slide_style_tables(self):
        stub = StubStreamlit(uploads={"cs_ak_upload": self._up()}, choices={"cs_source": "akamai"})
        self._run(stub)
        html = "".join(stub.html)
        for needle in ("Attack Type", "Total Hits", "By %", "Bad Reputation IP/Client"):
            self.assertIn(needle, html)

    def test_ips_snapshots_not_polluted_by_akamai(self):
        from cybersec_report.core import snapshot
        s2 = StubStreamlit(uploads={"cs_ak_upload": self._up()}, clicks={"cs_ak_save"}, choices={"cs_source": "akamai", "cs_ak_period_ok": True})
        self._run(s2)
        self.assertEqual(snapshot.list_snapshots(), [])

    def test_bad_file_shows_error_without_crash(self):
        up = FakeUpload(b"junk"); up.name = "x.xlsx"
        stub = StubStreamlit(uploads={"cs_ak_upload": up}, choices={"cs_source": "akamai"})
        self._run(stub)
        self.assertIn("error", stub.calls)
        self.assertNotIn("exception", stub.calls)


if __name__ == "__main__":
    unittest.main()
