"""
=========================================================
SOC Report Generator
Excel Writer v2.0
=========================================================
"""

from pathlib import Path
import pandas as pd

from openpyxl.styles import (
    PatternFill,
    Font,
    Border,
    Side,
    Alignment
)

from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

from core.charts import ChartBuilder


class ExcelWriter:

    def __init__(self, report, week):

        self.report = report
        self.week = week

        self.output_dir = Path("output")
        self.output_dir.mkdir(exist_ok=True)

        self.output_file = self.output_dir / f"SOC_Report_{week}.xlsx"

        # =============================
        # Theme
        # =============================

        self.header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        self.alt_fill = PatternFill(
            fill_type="solid",
            fgColor="F7F9FC"
        )

        self.header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        self.normal_font = Font(
            color="000000"
        )

        thin = Side(
            border_style="thin",
            color="D9D9D9"
        )

        self.border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin
        )

    # =====================================================
    # Auto Width
    # =====================================================

    def auto_width(self, ws):

        for column_cells in ws.columns:

            length = 0

            column = column_cells[0].column_letter

            for cell in column_cells:

                try:

                    if cell.value is not None:

                        length = max(
                            length,
                            len(str(cell.value))
                        )

                except:
                    pass

            width = min(max(length + 3, 12), 45)

            ws.column_dimensions[column].width = width

    # =====================================================
    # Format Worksheet
    # =====================================================

    def format_sheet(self, ws):

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions

        # Header
        for cell in ws[1]:

            cell.fill = self.header_fill

            cell.font = self.header_font

            cell.border = self.border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Data
        for row in ws.iter_rows(min_row=2):

            for cell in row:

                cell.border = self.border

                cell.font = self.normal_font

                cell.alignment = Alignment(
                    vertical="center"
                )

                # Zebra Row
                if cell.row % 2 == 0:
                    cell.fill = self.alt_fill

                # Number Formatting
                if isinstance(cell.value, float):

                    cell.number_format = "#,##0.00"

                elif isinstance(cell.value, int):

                    cell.number_format = "#,##0"

        self.auto_width(ws)

    # =====================================================
    # Write One Sheet
    # =====================================================

    def write_sheet(
        self,
        writer,
        sheet_name,
        df
    ):

        if not isinstance(df, pd.DataFrame):
            return

        sheet = sheet_name.replace("_", " ").title()

        sheet = sheet[:31]

        df.to_excel(

            writer,

            sheet_name=sheet,

            index=False

        )

        ws = writer.sheets[sheet]

        self.format_sheet(ws)

    # =====================================================
    # Export
    # =====================================================

    def export(self):

        print("\nWriting Excel Report...")

        with pd.ExcelWriter(
            self.output_file,
            engine="openpyxl"
        ) as writer:

            for sheet_name, df in self.report.items():

                self.write_sheet(
                    writer,
                    sheet_name,
                    df
                )

        print("Adding native Excel charts...")
        wb = load_workbook(self.output_file)
        ChartBuilder(wb).build_all()
        wb.save(self.output_file)

        print(f"\nReport Saved : {self.output_file}")

        return self.output_file