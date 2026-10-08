"""
=========================================================
Charts Generator (BARU - sebelumnya file ini kosong)
=========================================================
Membuat native Excel chart (bukan gambar statis) memakai openpyxl, supaya
tetap bisa diedit/diklik di Excel dan datanya nyambung ke sheet aslinya.
=========================================================
"""

from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList


class ChartBuilder:

    def __init__(self, workbook):
        self.wb = workbook

    # ---------------------------------------------------
    # Bar chart: Top 10 rule berdasarkan Incident_Count
    # ---------------------------------------------------
    def add_top_rule_chart(self, sheet_name="Top Rule", anchor="J2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        max_row = min(ws.max_row, 11)  # header + top 10

        chart = BarChart()
        chart.type = "col"
        chart.title = "Top 10 Rule - Incident Count"
        chart.y_axis.title = "Jumlah Incident"
        chart.x_axis.title = "Rule"
        chart.height, chart.width = 9, 20

        # kolom: Rule_ID, Rule_Names, Severity, Alert_Count, Incident_Count, ...
        data = Reference(ws, min_col=5, min_row=1, max_row=max_row)
        cats = Reference(ws, min_col=2, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Pie chart: distribusi Security Domain
    # ---------------------------------------------------
    def add_domain_pie(self, sheet_name="Domain", anchor="F2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        max_row = ws.max_row

        chart = PieChart()
        chart.title = "Distribusi Incident per Security Domain"
        chart.height, chart.width = 9, 14

        data = Reference(ws, min_col=3, min_row=1, max_row=max_row)  # Incident
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)  # Security_Domain
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Bar chart: workload per analyst
    # ---------------------------------------------------
    def add_soc_kpi_chart(self, sheet_name="Soc Individual Kpi", anchor="M2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        max_row = ws.max_row

        chart = BarChart()
        chart.type = "bar"
        chart.title = "Workload per Analyst (Total Incident)"
        chart.height, chart.width = 9, 20

        data = Reference(ws, min_col=2, min_row=1, max_row=max_row)  # Total_Incident
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)  # Owner
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Bar chart: coverage % alert non-Splunk per device
    # ---------------------------------------------------
    def add_nonsplunk_coverage_chart(self, sheet_name="Nonsplunk Coverage", anchor="H2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        if ws.max_row < 2:
            return

        chart = BarChart()
        chart.type = "col"
        chart.title = "Coverage Alert Non-Splunk (%) per Device"
        chart.y_axis.title = "% Jadi Tiket"
        chart.height, chart.width = 8, 16

        max_row = ws.max_row
        data = Reference(ws, min_col=6, min_row=1, max_row=max_row)  # Coverage_%
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)  # Source_Device
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Bar chart: pola alert per jam
    # ---------------------------------------------------
    def add_time_pattern_chart(self, sheet_name="Time By Hour", anchor="D2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        if ws.max_row < 2:
            return

        chart = BarChart()
        chart.type = "col"
        chart.title = "Distribusi Alert per Jam"
        chart.x_axis.title = "Jam"
        chart.y_axis.title = "Jumlah Alert"
        chart.height, chart.width = 8, 16

        max_row = ws.max_row
        data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Pie chart: breakdown insiden per severity
    # ---------------------------------------------------
    def add_severity_breakdown_chart(self, sheet_name="Severity Breakdown", anchor="F2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        if ws.max_row < 2:
            return

        chart = PieChart()
        chart.title = "Distribusi Tiket per Severity"
        chart.height, chart.width = 8, 14

        max_row = ws.max_row
        data = Reference(ws, min_col=2, min_row=1, max_row=max_row)   # Jumlah_Tiket
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)   # Severity_Label
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        ws.add_chart(chart, anchor)

    # ---------------------------------------------------
    # Line chart: tren harian alert & insiden
    # ---------------------------------------------------
    def add_daily_trend_chart(self, sheet_name="Daily Trend", anchor="E2"):
        if sheet_name not in self.wb.sheetnames:
            return
        ws = self.wb[sheet_name]
        if ws.max_row < 2:
            return

        chart = LineChart()
        chart.title = "Tren Harian: Alert vs Insiden"
        chart.x_axis.title = "Tanggal"
        chart.y_axis.title = "Jumlah"
        chart.height, chart.width = 9, 22
        chart.style = 2

        max_row = ws.max_row
        data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        for s in chart.series:
            s.smooth = False
        ws.add_chart(chart, anchor)

    def build_all(self):
        self.add_top_rule_chart()
        self.add_domain_pie()
        self.add_soc_kpi_chart()
        self.add_nonsplunk_coverage_chart()
        self.add_time_pattern_chart()
        self.add_severity_breakdown_chart()
        self.add_daily_trend_chart()
